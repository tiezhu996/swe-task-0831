#!/usr/bin/env python3
"""Create and validate a SWE-like delivery workbook."""

from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from check_prompt_independence import load_issues, validate_candidates


HEADERS = [
    "题目名称", "type", "提交人", "提交日期", "Repo URL", "Commit/版本", "主要语言",
    "任务类型", "需求 Prompt（原文）", "真实性与难度说明", "可能涉及模块",
    "Verify Rubric", "产物结果", "产物补充材料", "Seed 模型/版本",
    "Trae Session ID", "有效轮数", "是否完成需求", "Reviewer", "是否通过质检", "备注",
]
BASE_HEADERS = HEADERS
HEADERS = [
    "题目名称", "type", "提交人", "提交日期", "Repo URL", "Commit/版本", "Commit URL", "主要语言",
    "任务类型", "需求 Prompt（原文）", "真实性与难度说明", "可能涉及模块",
    "Verify Rubric", "产物结果", "产物补充材料", "Seed 模型/版本", "seed轮次",
    "Trae Session ID", "Trae Session ID 2", "有效轮数", "是否完成需求", "Reviewer", "是否通过质检", "备注",
]
LEGACY_HEADERS = [header for header in BASE_HEADERS if header != "type"]
OPTIONAL_HEADERS = {"type", "Reviewer", "是否通过质检", "Trae Session ID 2", "seed轮次"}
TASK_TYPE_OPTIONS = ("功能新增", "Bug 修复", "测试增强", "重构/性能", "配置/工具链", "其他", "问题修复")
RUN_TYPE_OPTIONS = ("有效轮数 > 100", "有效轮数 < 100 且 效果差", "有效轮数 < 100 且 效果好")
DELIVERY_STATUSES = ("未执行", "部分完成", "已完成")
RUBRIC_HEADER_PATTERN = re.compile(r"^\s*rubrics:\s*$", re.MULTILINE)
RUBRIC_NEW_ITEM_PATTERN = re.compile(r"^\s*-\s*id:\s*(\d+)\s*$", re.MULTILINE)
RUBRIC_NEW_TYPE_PATTERN = re.compile(r"^\s*type:\s*(f2p|p2p)\s*$")
RUBRIC_NEW_TEXT_PATTERN = re.compile(r"^\s*text:\s*(.+?)\s*$")
RUBRIC_LEGACY_PATTERN = re.compile(r"\[(f2p|p2p)\]\s*(\d+)\.")
RUBRIC_ITEM_PATTERN = re.compile(r"(?:^|；)\s*(\d+)\.\s*")
RESULT_LINE_PATTERN = re.compile(r"^[ ]*(\d+)[ ]+(通过|未通过)(?:[ ]+(.*))?$")
COMMIT_ID_PATTERN = re.compile(r"^[0-9a-f]{7,64}$", re.IGNORECASE)
COMMIT_ID_IN_TEXT_PATTERN = re.compile(r"[0-9a-f]{7,64}", re.IGNORECASE)
TRAE_SESSION_TOKEN_STRUCTURE_PATTERN = re.compile(r"^[^:\r\n]+:[^:\r\n]+:.+$")
TRAE_SESSION_ID_CANDIDATE_PATTERN = re.compile(r"(?<![0-9a-f])[0-9a-f]{24}(?![0-9a-f])", re.IGNORECASE)
GITHUB_REPOSITORY_URL_PATTERN = re.compile(
    r"^https?://(?:www\.)?github\.com/(?P<owner>[^/?#\s]+)/(?P<repo>[^/?#\s]+?)(?:\.git)?/?(?:[?#].*)?$",
    re.IGNORECASE,
)
GITHUB_COMMIT_URL_PATTERN = re.compile(
    r"^https?://(?:www\.)?github\.com/(?P<owner>[^/?#\s]+)/(?P<repo>[^/?#\s]+)/commit/(?P<sha>[0-9a-f]{7,64})(?:[?#].*)?$",
    re.IGNORECASE,
)
ERROR_MARKERS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NULL!", "#NUM!")
ISSUE_LEAK_PATTERNS = (
    re.compile(r"github\.com/[^\s/]+/[^\s/]+/issues?", re.IGNORECASE),
    re.compile(r"\bissue\s*#?\s*\d+", re.IGNORECASE),
    re.compile(r"参考\s*issue|公开问题|issue\s*编号", re.IGNORECASE),
)
WRAPPED_HEADERS = {"需求 Prompt（原文）", "真实性与难度说明", "可能涉及模块", "Verify Rubric", "产物结果", "产物补充材料", "备注"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--rows-json", required=True, help="JSON 文件，内容为对象数组")
    parser.add_argument("--output", required=True, help="输出 xlsx 路径")
    parser.add_argument("--template", help="可选的 SWE-like.xlsx 模板路径")
    parser.add_argument("--issues-json", required=True, help="本题研究时抓取的公开 Issue 标题/正文 JSON 数组；用于文案独立性硬校验")
    return parser.parse_args()


def load_rows(path: Path) -> list[dict]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, list) or not value:
        raise ValueError("rows-json 必须是非空对象数组")
    for index, row in enumerate(value, start=1):
        if not isinstance(row, dict):
            raise ValueError(f"第 {index} 行不是对象")
    return value


def insert_header_column(ws, index: int, header: str) -> None:
    ws.insert_cols(index)
    source_column = min(index + 1, ws.max_column)
    for row in range(1, ws.max_row + 1):
        source = ws.cell(row, source_column)
        target = ws.cell(row, index)
        target._style = copy.copy(source._style)
        target.font = copy.copy(source.font)
        target.alignment = copy.copy(source.alignment)
    source_letter = get_column_letter(source_column)
    target_letter = get_column_letter(index)
    ws.column_dimensions[target_letter].width = ws.column_dimensions[source_letter].width
    ws.cell(1, index).value = header


def upgrade_template_headers(ws) -> list[str]:
    headers = [cell.value for cell in ws[1]]
    if headers[: len(LEGACY_HEADERS)] == LEGACY_HEADERS:
        insert_header_column(ws, 2, "type")
        headers = [cell.value for cell in ws[1]]
    if headers[: len(BASE_HEADERS)] == BASE_HEADERS:
        insert_header_column(ws, 7, "Commit URL")
        headers = [cell.value for cell in ws[1]]
        insert_header_column(ws, headers.index("Seed 模型/版本") + 2, "seed轮次")
        headers = [cell.value for cell in ws[1]]
        insert_header_column(ws, headers.index("Trae Session ID") + 2, "Trae Session ID 2")
        headers = [cell.value for cell in ws[1]]
    return headers


def create_workbook(template: Path | None):
    if template:
        wb = load_workbook(template)
        if "数据表" not in wb.sheetnames:
            raise ValueError("模板中缺少工作表：数据表")
        ws = wb["数据表"]
        headers = upgrade_template_headers(ws)
        if headers[: len(HEADERS)] != HEADERS:
            raise ValueError("模板表头与 SWE-like 24 列不一致")
        style_source = [copy.copy(ws.cell(2, col)._style) for col in range(1, len(HEADERS) + 1)] if ws.max_row >= 2 else None
        font_source = [copy.copy(ws.cell(2, col).font) for col in range(1, len(HEADERS) + 1)] if ws.max_row >= 2 else None
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        return wb, ws, style_source, font_source

    wb = Workbook()
    ws = wb.active
    ws.title = "数据表"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10)
    for column in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(column)].width = 19
    return wb, ws, None, None


def value_for(row: dict, header: str):
    value = row.get(header)
    if value is None or value == "":
        if header in OPTIONAL_HEADERS:
            return ""
        raise ValueError(f"缺少或为空字段：{header}")
    if header == "提交日期" and isinstance(value, str):
        try:
            return datetime.fromisoformat(value)
        except ValueError as exc:
            raise ValueError(f"提交日期不是 ISO 日期：{value}") from exc
    return value


def _finalize_rubric_item(current: dict) -> tuple[int, str, str]:
    identifier = current["id"]
    item_type = current.get("type")
    text = current.get("text")
    if item_type not in ("f2p", "p2p"):
        raise ValueError(f"Verify Rubric 第 {identifier} 条缺少或非法 type，只能填 f2p 或 p2p")
    if not text:
        raise ValueError(f"Verify Rubric 第 {identifier} 条缺少 text")
    return identifier, item_type, text


def parse_rubric_items(rubric: str) -> list[tuple[int, str, str]]:
    value = str(rubric).strip()
    if not value:
        raise ValueError("Verify Rubric 不能为空")

    if RUBRIC_NEW_ITEM_PATTERN.search(value) or RUBRIC_HEADER_PATTERN.search(value):
        items: list[tuple[int, str, str]] = []
        current: dict | None = None
        for raw in value.splitlines():
            line = raw.strip()
            if not line:
                continue
            if RUBRIC_HEADER_PATTERN.fullmatch(line):
                continue
            match = RUBRIC_NEW_ITEM_PATTERN.fullmatch(line)
            if match:
                if current is not None:
                    items.append(_finalize_rubric_item(current))
                current = {"id": int(match.group(1)), "type": None, "text": None}
                continue
            if current is not None:
                type_match = RUBRIC_NEW_TYPE_PATTERN.fullmatch(line)
                if type_match:
                    if current["type"] is not None:
                        raise ValueError(f"Verify Rubric 第 {current['id']} 条重复填写 type")
                    current["type"] = type_match.group(1)
                    continue
                text_match = RUBRIC_NEW_TEXT_PATTERN.fullmatch(line)
                if text_match:
                    if current["text"] is not None:
                        raise ValueError(f"Verify Rubric 第 {current['id']} 条重复填写 text")
                    current["text"] = text_match.group(1)
                    continue
            raise ValueError("Verify Rubric 必须按 rubrics 书写范例填写：- id: 1 / type: f2p / text: ...")
        if current is not None:
            items.append(_finalize_rubric_item(current))
        if not items:
            raise ValueError("Verify Rubric 至少包含 1 条")
        return items

    legacy = RUBRIC_LEGACY_PATTERN.findall(value)
    if legacy:
        return [(int(number), item_type, "") for item_type, number in legacy]
    numbers = [int(item) for item in RUBRIC_ITEM_PATTERN.findall(value)]
    if numbers:
        return [(number, "f2p", "") for number in numbers]
    raise ValueError("Verify Rubric 必须使用 rubrics 结构或旧版连续编号格式")


def parse_result_lines(result: str) -> list[tuple[int, str, str]]:
    items: list[tuple[int, str, str]] = []
    seen: set[int] = set()
    for raw in str(result).splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        match = RESULT_LINE_PATTERN.fullmatch(line)
        if not match:
            raise ValueError("产物结果每行必须是：编号 + 普通空格 + 通过/未通过，不能带句点、顿号、冒号、括号或制表符")
        number = int(match.group(1))
        conclusion = match.group(2)
        reason = (match.group(3) or "").strip()
        if conclusion == "未通过" and not reason:
            raise ValueError(f"产物结果第 {number} 条为未通过时必须写原因")
        if conclusion == "通过" and reason:
            raise ValueError(f"产物结果第 {number} 条为通过时不能附带原因")
        if number in seen:
            raise ValueError(f"产物结果编号 {number} 重复")
        seen.add(number)
        items.append((number, conclusion, reason))
    return items


def validate_row_contract(row: dict) -> None:
    for header in ("Trae Session ID 2", "seed轮次"):
        if row.get(header) not in (None, ""):
            raise ValueError(f"{header} 是新版占位字段，必须留空")
    trae_session = row.get("Trae Session ID", "")
    if trae_session != "未提供":
        if not isinstance(trae_session, str) or trae_session != trae_session.strip():
            raise ValueError("Trae Session ID 必须填写完整原始 Trae Session token，不能有首尾截断或空白")
        if "..." in trae_session or "…" in trae_session:
            raise ValueError("Trae Session ID 不得使用省略号；必须填写完整原始 Trae Session token")
        if not TRAE_SESSION_TOKEN_STRUCTURE_PATTERN.fullmatch(trae_session) or not TRAE_SESSION_ID_CANDIDATE_PATTERN.search(trae_session):
            raise ValueError(
                "Trae Session ID 必须填写完整原始 Trae Session token；"
                "不得只填 24 位 chat_session_id。没有完整 token 时填写“未提供”"
            )
    task_type = row.get("任务类型")
    if task_type not in TASK_TYPE_OPTIONS:
        raise ValueError(f"任务类型必须是以下之一：{', '.join(TASK_TYPE_OPTIONS)}")
    run_type = row.get("type") or ""
    if run_type not in ("", *RUN_TYPE_OPTIONS):
        raise ValueError(f"type 必须是以下之一或留空：{', '.join(RUN_TYPE_OPTIONS)}")
    rounds = row.get("有效轮数")
    if run_type == "有效轮数 > 100" and (not isinstance(rounds, int) or rounds <= 100):
        raise ValueError("type 为“有效轮数 > 100”时，有效轮数必须是大于 100 的整数")
    if run_type in ("有效轮数 < 100 且 效果差", "有效轮数 < 100 且 效果好") and (not isinstance(rounds, int) or rounds >= 100):
        raise ValueError("type 为“有效轮数 < 100”时，有效轮数必须是小于 100 的整数")
    if isinstance(rounds, int) and rounds == 100 and run_type:
        raise ValueError("有效轮数恰为 100 时没有匹配的 type 选项，type 必须留空并在备注中待人工确认")
    completion = row.get("是否完成需求")
    if completion not in DELIVERY_STATUSES:
        raise ValueError(f"是否完成需求必须是以下之一：{', '.join(DELIVERY_STATUSES)}")
    rubric_items = parse_rubric_items(row.get("Verify Rubric", ""))
    numbers = [item[0] for item in rubric_items]
    if len(numbers) < 5:
        raise ValueError("Verify Rubric 必须至少包含 5 项连续编号")
    if numbers != list(range(1, len(numbers) + 1)):
        raise ValueError("Verify Rubric 的 id 必须从 1 连续递增，不重号、不跳号")
    if not any(item[1] == "f2p" for item in rubric_items):
        raise ValueError("Verify Rubric 必须至少包含 1 条 f2p")
    result = row.get("产物结果", "")
    if completion == "未执行":
        if not result.startswith("未执行"):
            raise ValueError("未执行题目的产物结果必须以“未执行”起句并说明尚无逐条质检结果")
    else:
        result_items = parse_result_lines(result)
        result_numbers = [item[0] for item in result_items]
        if result_numbers != numbers:
            raise ValueError("产物结果必须逐条对应 Verify Rubric 的 id：通过写 N 通过，未通过写 N 未通过 原因")
    revision = row.get("Commit/版本", "")
    if not isinstance(revision, str) or not revision.strip():
        raise ValueError("Commit/版本 必须填写单独的 Commit ID 或版本号")
    if COMMIT_ID_IN_TEXT_PATTERN.search(revision) and not COMMIT_ID_PATTERN.fullmatch(revision.strip()):
        raise ValueError("Commit/版本 只能填写 Commit ID 或版本号之一；存在 Commit ID 时请只填写 Commit ID")
    source_url = row.get("Repo URL", "")
    source_match = GITHUB_REPOSITORY_URL_PATTERN.fullmatch(source_url.strip()) if isinstance(source_url, str) else None
    if not source_match:
        raise ValueError("Repo URL 必须是原始 GitHub 仓库根地址，例如 https://github.com/owner/repo")
    commit_url = row.get("Commit URL", "")
    commit_match = GITHUB_COMMIT_URL_PATTERN.fullmatch(commit_url.strip()) if isinstance(commit_url, str) else None
    if not commit_match:
        raise ValueError("Commit URL 必须是 fork 仓库中任务变更提交的 GitHub commit URL")
    source_repo = (source_match.group("owner").lower(), source_match.group("repo").removesuffix(".git").lower())
    fork_repo = (commit_match.group("owner").lower(), commit_match.group("repo").lower())
    if source_repo == fork_repo:
        raise ValueError("Commit URL 必须指向自有 fork 仓库的任务提交，不能与 Repo URL 指向同一仓库")
    if COMMIT_ID_PATTERN.fullmatch(revision.strip()) and revision.lower() == commit_match.group("sha").lower():
        raise ValueError("Commit URL 必须指向任务变更后的新提交，不能与原始 Commit/版本 相同")


def append_rows(ws, rows, style_source, font_source):
    for row in rows:
        validate_row_contract(row)
        ws.append([value_for(row, header) for header in HEADERS])
        row_index = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row_index, col)
            if style_source:
                cell._style = copy.copy(style_source[col - 1])
            if font_source:
                cell.font = copy.copy(font_source[col - 1])
            if HEADERS[col - 1] in WRAPPED_HEADERS:
                cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
        for header in ("Repo URL", "Commit URL"):
            url_cell = ws.cell(row_index, HEADERS.index(header) + 1)
            if isinstance(url_cell.value, str) and url_cell.value.startswith("http"):
                url_cell.hyperlink = url_cell.value
                url_cell.style = "Hyperlink"
        ws.cell(row_index, HEADERS.index("提交日期") + 1).number_format = "yyyy-mm-dd"
        ws.row_dimensions[row_index].height = 150


def validate_issue_independence(rows: list[dict], issues_json: Path) -> None:
    issues = load_issues(issues_json)
    candidates: dict[str, str] = {}
    protected_fields = ("需求 Prompt（原文）", "真实性与难度说明", "Verify Rubric")
    for index, row in enumerate(rows, start=1):
        for field in protected_fields:
            candidates[f"第 {index} 行 {field}"] = str(row.get(field, ""))
    report = validate_candidates(candidates, issues)
    if report["status"] == "blocked":
        print(json.dumps(report, ensure_ascii=False, indent=2), file=sys.stderr)
        raise ValueError("Issue 文案独立性校验未通过，禁止生成交付表；请重新设计题面、真实性说明或 Verify Rubric 后再试")


def configure_dropdowns(ws) -> None:
    ws.data_validations.dataValidation = []
    last_row = max(ws.max_row, 2)
    for header, options, allow_blank in (
        ("type", RUN_TYPE_OPTIONS, True),
        ("任务类型", TASK_TYPE_OPTIONS, False),
    ):
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(options) + '"',
            allow_blank=allow_blank,
        )
        validation.error = f"请选择：{'、'.join(options)}"
        validation.errorTitle = f"{header} 取值无效"
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        column = get_column_letter(HEADERS.index(header) + 1)
        validation.add(f"{column}2:{column}{last_row}")


def validate_workbook(path: Path) -> None:
    for data_only in (False, True):
        wb = load_workbook(path, data_only=data_only)
        try:
            ws = wb["数据表"]
            headers = [cell.value for cell in ws[1]]
            if headers[: len(HEADERS)] != HEADERS:
                raise ValueError("输出表头不完整")
            for row in range(2, ws.max_row + 1):
                for col in range(1, len(HEADERS) + 1):
                    value = ws.cell(row, col).value
                    if (value is None or value == "") and headers[col - 1] not in OPTIONAL_HEADERS:
                        raise ValueError(f"{ws.cell(row, col).coordinate} 为空")
                    if isinstance(value, str) and any(marker in value for marker in ERROR_MARKERS):
                        raise ValueError(f"{ws.cell(row, col).coordinate} 包含公式错误：{value}")
                    if isinstance(value, str) and any(pattern.search(value) for pattern in ISSUE_LEAK_PATTERNS):
                        raise ValueError(f"{ws.cell(row, col).coordinate} 暴露了 Issue 来源信息")
        finally:
            wb.close()


def main() -> None:
    args = parse_args()
    rows = load_rows(Path(args.rows_json))
    validate_issue_independence(rows, Path(args.issues_json))
    template = Path(args.template) if args.template else None
    wb, ws, style_source, font_source = create_workbook(template)
    append_rows(ws, rows, style_source, font_source)
    configure_dropdowns(ws)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    validate_workbook(output)
    print(json.dumps({"status": "success", "rows": len(rows), "output": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
