import subprocess
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parents[1] / "scripts"))
from build_delivery_xlsx import (
    BASE_HEADERS,
    HEADERS,
    append_rows,
    configure_dropdowns,
    create_workbook,
    validate_row_contract,
    validate_workbook,
)
from publish_task_commits import apply_task_changes, repository_from_url


FULL_TRAE_SESSION_TOKEN = (
    "573312257230252:2e67421c5031754536c336fad2bf684e_"
    "6a982767a61b7e43f2ffbd69.6a982767a61b7e43f2ffbd6c.6a982767a61b7e43f2ffbd6a:"
    "TraeCode CN.3.3.90.no_sid.no_ppe.T(2026/9/2 21:40:55)"
)


def valid_row():
    return {
        "题目名称": "测试题",
        "type": "",
        "提交人": "tester",
        "提交日期": date(2026, 9, 3),
        "Repo URL": "https://github.com/upstream/project",
        "Commit/版本": "0123456789abcdef",
        "Commit URL": "https://github.com/example/project-fork/commit/abcdef0123456789",
        "主要语言": "Python",
        "任务类型": "功能新增",
        "需求 Prompt（原文）": "请改善任务处理体验。",
        "真实性与难度说明": "需要同时处理成功、失败和重复执行后的可观察状态。",
        "可能涉及模块": "任务调度、状态存储和测试模块",
        "Verify Rubric": "rubrics:\n- id: 1\n  type: f2p\n  text: 默认状态可保持兼容。\n- id: 2\n  type: f2p\n  text: 主流程产生可观察结果。\n- id: 3\n  type: f2p\n  text: 失败后保留可恢复状态。\n- id: 4\n  type: f2p\n  text: 重复请求不会重复写入。\n- id: 5\n  type: f2p\n  text: 持久化边界可验证。",
        "产物结果": "1 通过\n2 通过\n3 通过\n4 通过\n5 通过",
        "产物补充材料": "已执行；测试结果可复核。",
        "Seed 模型/版本": "未提供",
        "seed轮次": "",
        "Trae Session ID": "未提供",
        "Trae Session ID 2": "",
        "有效轮数": "未提供",
        "是否完成需求": "已完成",
        "Reviewer": "",
        "是否通过质检": "",
        "备注": "已完成。",
    }


class BuildDeliveryWorkbookTests(unittest.TestCase):
    def test_contract_accepts_empty_new_tracking_fields(self):
        validate_row_contract(valid_row())

    def test_contract_rejects_source_repository_commit_url(self):
        row = valid_row()
        row["Commit URL"] = "https://github.com/upstream/project/commit/abcdef0123456789"
        with self.assertRaisesRegex(ValueError, "自有 fork"):
            validate_row_contract(row)

    def test_contract_rejects_nonempty_new_tracking_fields(self):
        row = valid_row()
        row["seed轮次"] = "1"
        with self.assertRaisesRegex(ValueError, "必须留空"):
            validate_row_contract(row)

    def test_contract_accepts_full_trae_session_token(self):
        row = valid_row()
        row["Trae Session ID"] = FULL_TRAE_SESSION_TOKEN
        validate_row_contract(row)

    def test_contract_rejects_short_or_truncated_trae_session_id(self):
        row = valid_row()
        row["Trae Session ID"] = "6a982767a61b7e43f2ffbd69"
        with self.assertRaisesRegex(ValueError, "完整原始"):
            validate_row_contract(row)
        row["Trae Session ID"] = f"{FULL_TRAE_SESSION_TOKEN[:60]}..."
        with self.assertRaisesRegex(ValueError, "省略号"):
            validate_row_contract(row)

    def test_contract_rejects_dot_numbered_result(self):
        row = valid_row()
        row["产物结果"] = "1. 通过\n2. 通过\n3. 通过\n4. 通过\n5. 通过"
        with self.assertRaisesRegex(ValueError, "普通空格"):
            validate_row_contract(row)

    def test_contract_rejects_failed_item_without_reason(self):
        row = valid_row()
        row["产物结果"] = "1 通过\n2 未通过\n3 通过\n4 通过\n5 通过"
        with self.assertRaisesRegex(ValueError, "原因"):
            validate_row_contract(row)

    def test_contract_rejects_duplicate_result_id(self):
        row = valid_row()
        row["产物结果"] = "1 通过\n1 通过\n2 通过\n3 通过\n4 通过"
        with self.assertRaisesRegex(ValueError, "重复"):
            validate_row_contract(row)

    def test_contract_rejects_missing_result_id(self):
        row = valid_row()
        row["产物结果"] = "1 通过\n2 通过\n4 通过\n5 通过"
        with self.assertRaisesRegex(ValueError, "逐条对应"):
            validate_row_contract(row)

    def test_old_template_is_upgraded_to_current_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "legacy.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "数据表"
            worksheet.append(BASE_HEADERS)
            worksheet.append([""] * len(BASE_HEADERS))
            workbook.save(template)
            workbook.close()

            workbook, worksheet, _, _ = create_workbook(template)
            self.assertEqual([cell.value for cell in worksheet[1]][: len(HEADERS)], HEADERS)
            workbook.close()

    def test_legacy_twenty_column_template_is_upgraded_to_current_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "legacy.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "数据表"
            worksheet.append([header for header in BASE_HEADERS if header != "type"])
            worksheet.append([""] * (len(BASE_HEADERS) - 1))
            workbook.save(template)
            workbook.close()

            workbook, worksheet, _, _ = create_workbook(template)
            self.assertEqual([cell.value for cell in worksheet[1]][: len(HEADERS)], HEADERS)
            workbook.close()

    def test_generated_workbook_has_current_headers_and_url_links(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "delivery.xlsx"
            workbook, worksheet, styles, fonts = create_workbook(None)
            append_rows(worksheet, [valid_row()], styles, fonts)
            configure_dropdowns(worksheet)
            workbook.save(output)
            workbook.close()

            validate_workbook(output)
            from openpyxl import load_workbook

            workbook = load_workbook(output)
            worksheet = workbook["数据表"]
            self.assertEqual([cell.value for cell in worksheet[1]], HEADERS)
            self.assertEqual(worksheet.cell(2, HEADERS.index("Repo URL") + 1).hyperlink.target, valid_row()["Repo URL"])
            self.assertEqual(worksheet.cell(2, HEADERS.index("Commit URL") + 1).hyperlink.target, valid_row()["Commit URL"])
            workbook.close()

    def test_commit_publisher_parses_github_repository_root_url(self):
        repository = repository_from_url("https://github.com/upstream/project.git", "repo-url")
        self.assertEqual(repository.slug, "upstream/project")
        self.assertEqual(repository.url, "https://github.com/upstream/project")

    def test_publisher_applies_only_tracked_and_unignored_task_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source"
            destination = Path(directory) / "destination"
            source.mkdir()
            for command in (
                ["git", "init", "--quiet", str(source)],
                ["git", "-C", str(source), "config", "user.name", "tester"],
                ["git", "-C", str(source), "config", "user.email", "tester@example.com"],
            ):
                subprocess.run(command, check=True)
            (source / ".gitignore").write_text("build/\n", encoding="utf-8")
            (source / "tracked.txt").write_text("before\n", encoding="utf-8")
            subprocess.run(["git", "-C", str(source), "add", "."], check=True)
            subprocess.run(["git", "-C", str(source), "commit", "--quiet", "-m", "base"], check=True)
            baseline = subprocess.run(
                ["git", "-C", str(source), "rev-parse", "HEAD"], text=True, capture_output=True, check=True
            ).stdout.strip()
            subprocess.run(["git", "clone", "--quiet", str(source), str(destination)], check=True)
            (source / "tracked.txt").write_text("after\n", encoding="utf-8")
            (source / "new.txt").write_text("new\n", encoding="utf-8")
            (source / "build").mkdir()
            (source / "build" / "cache.bin").write_text("ignored\n", encoding="utf-8")

            apply_task_changes(source, destination, baseline, allow_unverified=False)

            self.assertEqual((destination / "tracked.txt").read_text(encoding="utf-8"), "after\n")
            self.assertEqual((destination / "new.txt").read_text(encoding="utf-8"), "new\n")
            self.assertFalse((destination / "build" / "cache.bin").exists())


if __name__ == "__main__":
    unittest.main()
